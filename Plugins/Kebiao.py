"""课表查询：导入教务课表、按周出图、上课提醒。

只做两件事：**查课表** 和 **上课提醒**（默认关闭，用指令开启）。
所有功能都走指令触发，不监听任何关键字，也不会在群里主动播报。

用法：

- 把教务系统导出的「学生个人课表_*.xls / *.xlsx」发到群里或私聊机器人，自动解析保存；
- ``/课表`` 看本周、``/周课表 8`` 看第 8 周、``/今日课表`` / ``/明日课表``、``/下一节课``；
- ``/上课提醒 开`` 打开课前 20 分钟提醒（默认关闭），``/上课提醒 关`` 关掉。
"""

import asyncio
import hashlib
import io
import json
import re
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from httpx import AsyncClient
from nonebot import get_driver, on_command, on_message, on_notice
from nonebot.adapters.onebot.v11 import (
    Bot,
    GroupUploadNoticeEvent,
    Message,
    MessageEvent,
    MessageSegment,
)
from nonebot.params import CommandArg
from nonebot.plugin import PluginMetadata
from PIL import Image, ImageDraw, ImageFont

from .HelpEntry import register_help_commands

# ===== 课表基础 =====

WEEKDAYS = ['周一', '周二', '周三', '周四', '周五', '周六', '周日']
WEEKDAY_KEYS = ['星期一', '星期二', '星期三', '星期四', '星期五', '星期六', '星期日']

# 教务导出的左列文案 -> 大节
PERIOD_LABEL = {
    '第一二节': '1-2',
    '第三四节': '3-4',
    '第五六节': '5-6',
    '第七八节': '7-8',
    '第九十节': '9-10',
    '第十一十二节': '11-12',
}

# 一条课程：名字 / 老师 / 3-12([周])[01-02节] / 地点
ENTRY_RE = re.compile(
    r'(?P<name>[^\n]+)\n'
    r'(?P<teacher>[^\n]+)\n'
    r'(?P<weeks>[0-9,\-]+)\(\[周\]\)\[(?P<period>[^\]]*)\]'
    r'(?:\n(?P<loc>[^\n]*))?'
)

# 第 1 周的周一，可用 /课表开学 2026-09-07 修改
DEFAULT_TERM_START = '2026-09-07'

# 课前多少分钟提醒
REMIND_MIN = 20
# 提醒循环的检查间隔（秒）
REMIND_TICK = 30

# 上传文件体积上限
FILE_MAX_BYTES = 5 * 1024 * 1024

# 数据目录：一个 QQ 一个 json
DATA_DIR = Path('Data/Kebiao')

# ===== 出图 =====

BLOCKS = [1, 3, 5, 7, 9]  # 大节起点
BLOCK_LABEL = {1: '1-2节', 3: '3-4节', 5: '5-6节', 7: '7-8节', 9: '9-10节'}
BLOCK_TIME = {1: '08:00', 3: '10:00', 5: '14:00', 7: '16:00', 9: '19:00'}

PALETTE = [
    (232, 241, 255),
    (255, 240, 232),
    (233, 248, 236),
    (255, 246, 224),
    (243, 236, 253),
    (230, 247, 249),
    (252, 236, 243),
    (238, 243, 228),
]

# 水印文案（作者署名，做得淡一点）
WATERMARK = '一香芹酮'

FONT_CANDIDATES = [
    'Data/Kebiao/NotoSansSC.ttf',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc',
    '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',
    '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
    'C:/Windows/Fonts/msyh.ttc',
    'C:/Windows/Fonts/simhei.ttf',
    '/System/Library/Fonts/PingFang.ttc',
]
BOLD_CANDIDATES = [
    'Data/Kebiao/NotoSansSC.ttf',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc',
    'C:/Windows/Fonts/msyhbd.ttc',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
    'C:/Windows/Fonts/msyh.ttc',
]
# 系统里一个中文字体都没有时，自动下载一份（约 10MB，只下一次）
FONT_URLS = (
    'https://fonts.gstatic.com/s/notosanssc/v41/k3kCo84MPvpLmixcA63oeAL7Iqp5IZJF9bmaG-3FnYw.ttf',
    'https://cdn.jsdelivr.net/gh/googlefonts/noto-cjk@main/Sans/SubsetOTF/SC/NotoSansSC-Regular.otf',
)

_FONT_PATH: Path | None = None
_FONT_CACHE: dict[tuple[int, bool], Any] = {}


def _find_font() -> Path | None:
    """返回系统里能用的字体文件（没有就返回 None）。"""
    for name in FONT_CANDIDATES:
        path = Path(name)
        if path.exists():
            return path
    return None


async def ensure_font() -> Path | None:
    """确保有中文字体可用：本地没有时下载一份到 Data/Kebiao。"""
    global _FONT_PATH
    if _FONT_PATH is not None:
        return _FONT_PATH
    _FONT_PATH = _find_font()
    if _FONT_PATH is not None:
        return _FONT_PATH

    target = DATA_DIR / 'NotoSansSC.ttf'
    target.parent.mkdir(parents=True, exist_ok=True)
    async with AsyncClient(timeout=60) as client:
        for url in FONT_URLS:
            try:
                response = await client.get(url, follow_redirects=True)
                if response.status_code != 200 or len(response.content) < 1024 * 1024:
                    continue
                target.write_bytes(response.content)
                _FONT_PATH = target
                return _FONT_PATH
            except Exception:
                continue
    return None


def _font(size: int, bold: bool = False):
    """按字号取字体；没有字体文件时退回 Pillow 内置字体。"""
    key = (size, bold)
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    font = None
    for name in BOLD_CANDIDATES if bold else FONT_CANDIDATES:
        path = Path(name)
        if not path.exists():
            continue
        try:
            font = ImageFont.truetype(str(path), size)
            break
        except Exception:
            continue
    if font is None:
        try:
            font = ImageFont.load_default(size)
        except TypeError:
            font = ImageFont.load_default()
    _FONT_CACHE[key] = font
    return font


def _color(name: str) -> tuple[int, int, int]:
    return PALETTE[int(hashlib.md5(name.encode('utf-8')).hexdigest()[:4], 16) % len(PALETTE)]


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, max_w: int) -> list[str]:
    lines, cur = [], ''
    for char in text:
        if draw.textlength(cur + char, font=font) <= max_w:
            cur += char
        else:
            lines.append(cur)
            cur = char
    if cur:
        lines.append(cur)
    return lines or ['']


def _watermark(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    """右下角一枚很淡的署名。"""
    font = _font(13)
    text_w = draw.textlength(WATERMARK, font=font)
    draw.text((width - text_w - 16, height - 22), WATERMARK, font=font, fill=(207, 213, 224))


def _merge(courses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """同一时段、同一门课的多条记录（不同周次 / 教室）合并成一张卡片。"""
    merged: dict[tuple, dict[str, Any]] = {}
    order: list[tuple] = []
    for course in courses:
        key = (course.get('day'), course.get('period'), course.get('name'))
        if key not in merged:
            item = dict(course)
            item['_locs'] = [course.get('location')] if course.get('location') else []
            merged[key] = item
            order.append(key)
            continue
        item = merged[key]
        item['weeks'] = sorted(set(item.get('weeks') or []) | set(course.get('weeks') or []))
        location = course.get('location')
        if location and location not in item['_locs']:
            item['_locs'].append(location)
    result = []
    for key in order:
        item = merged[key]
        locations = [loc for loc in item.get('_locs', []) if loc]
        item['location'] = ' / '.join(locations[:2]) + ('…' if len(locations) > 2 else '')
        result.append(item)
    return result


def _fmt_weeks(weeks) -> str:
    weeks = sorted({int(week) for week in weeks or []})
    if not weeks:
        return ''
    out, start, prev = [], weeks[0], weeks[0]
    for week in weeks[1:]:
        if week == prev + 1:
            prev = week
            continue
        out.append(f'{start}-{prev}' if start != prev else f'{start}')
        start = prev = week
    out.append(f'{start}-{prev}' if start != prev else f'{start}')
    return '、'.join(out) + '周'


def _block_of(number: int) -> int:
    """把 1~12 节映射到大节起点 1/3/5/7/9。"""
    if number <= 2:
        return 1
    if number <= 4:
        return 3
    if number <= 6:
        return 5
    if number <= 8:
        return 7
    return 9


def _span(period: str) -> tuple[int, int]:
    """'05-06节' -> (5, 5) 大节区间。"""
    numbers = [int(item) for item in re.findall(r'\d+', str(period))]
    if not numbers:
        return 1, 1
    return _block_of(numbers[0]), _block_of(numbers[-1])


def render_grid(
    courses: list[dict[str, Any]],
    title: str,
    subtitle: str = '',
    show_weeks: bool = False,
    footer: str = '',
    highlight: str | None = None,
    day_dates: dict[str, str] | None = None,
) -> bytes:
    """把课表画成整周网格图，返回 PNG 字节。"""
    col_w, row_h = 196, 150
    left_w, top_h = 122, 128
    head_h = 54 if day_dates else 44
    grid_top = top_h + head_h
    width = left_w + col_w * 7
    height = grid_top + row_h * len(BLOCKS) + 40
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)

    f_title = _font(40, True)
    f_sub = _font(24)
    f_head = _font(26, True)
    f_label = _font(22, True)
    f_tiny = _font(15)
    f_date = _font(14)

    # 标题与副标题
    draw.text((left_w // 2, 24), title, font=f_title, fill=(28, 32, 40))
    if subtitle:
        draw.text((left_w // 2, 78), subtitle, font=f_sub, fill=(110, 116, 128))

    # 表头：高亮那一天用蓝底
    for index, (label, day_key) in enumerate(zip(WEEKDAYS, WEEKDAY_KEYS)):
        x = left_w + index * col_w
        hot = day_key == highlight
        draw.rectangle([x, top_h, x + col_w, grid_top], fill=(64, 106, 214) if hot else (238, 242, 248))
        ink = (255, 255, 255) if hot else (40, 46, 58)
        date_text = (day_dates or {}).get(day_key, '')
        label_w = draw.textlength(label, font=f_head)
        if date_text:
            draw.text((x + (col_w - label_w) / 2, top_h + 4), label, font=f_head, fill=ink)
            date_w = draw.textlength(date_text, font=f_date)
            draw.text((x + (col_w - date_w) / 2, top_h + 34), date_text, font=f_date, fill=ink)
        else:
            draw.text((x + (col_w - label_w) / 2, top_h + 9), label, font=f_head, fill=ink)

    # 左侧节次与隔行底色
    for index, base in enumerate(BLOCKS):
        y = grid_top + index * row_h
        if index % 2 == 1:
            draw.rectangle([0, y, width, y + row_h], fill=(250, 251, 253))
        draw.text((14, y + 32), BLOCK_LABEL[base], font=f_label, fill=(60, 66, 78))
        draw.text((14, y + 68), BLOCK_TIME[base], font=f_tiny, fill=(140, 146, 158))

    # 网格线
    for index in range(len(BLOCKS) + 1):
        y = grid_top + index * row_h
        draw.line([(0, y), (width, y)], fill=(222, 227, 235), width=1)
    for index in range(8):
        x = left_w + index * col_w
        draw.line([(x, top_h), (x, grid_top + len(BLOCKS) * row_h)], fill=(222, 227, 235), width=1)
    draw.line([(0, top_h), (width, top_h)], fill=(222, 227, 235), width=1)

    # 高亮那列两侧描边
    if highlight and highlight in WEEKDAY_KEYS:
        hotspot = left_w + WEEKDAY_KEYS.index(highlight) * col_w
        bottom = grid_top + len(BLOCKS) * row_h
        draw.rectangle([hotspot, top_h, hotspot + 3, bottom], fill=(64, 106, 214))
        draw.rectangle([hotspot + col_w - 3, top_h, hotspot + col_w, bottom], fill=(64, 106, 214))

    # 同一格里的多门课 -> 分道并排
    groups: dict[tuple[int, int, int], list[dict[str, Any]]] = {}
    for course in _merge(courses):
        if course.get('day') not in WEEKDAY_KEYS:
            continue
        day_index = WEEKDAY_KEYS.index(course['day'])
        start_block, end_block = _span(course.get('period', ''))
        if start_block not in BLOCKS or end_block not in BLOCKS:
            continue
        groups.setdefault((day_index, BLOCKS.index(start_block), BLOCKS.index(end_block)), []).append(course)

    for (day_index, row_start, row_end), items in groups.items():
        count = len(items)
        base_x = left_w + day_index * col_w
        lane_w = (col_w - 8) / count
        y0 = grid_top + row_start * row_h + 4
        y1 = grid_top + (row_end + 1) * row_h - 4
        f_name = _font(17) if count == 1 else _font(16)
        for index, course in enumerate(items):
            x0 = base_x + 4 + index * lane_w
            x1 = x0 + lane_w - 3
            draw.rounded_rectangle([x0, y0, x1, y1], radius=8, fill=_color(course['name']))
            padding = 6
            inner = int(x1 - x0 - padding * 2)
            text_y = y0 + padding
            for line in _wrap(draw, course.get('name', ''), f_name, inner)[:5]:
                draw.text((x0 + padding, text_y), line, font=f_name, fill=(30, 36, 46))
                text_y += 20
            for line in _wrap(draw, course.get('location', ''), f_tiny, inner)[:3]:
                draw.text((x0 + padding, text_y + 1), line, font=f_tiny, fill=(88, 96, 110))
                text_y += 18
            if show_weeks:
                weeks_text = _fmt_weeks(course.get('weeks'))
                if weeks_text:
                    week_lines = _wrap(draw, weeks_text, f_tiny, inner)[:2]
                    week_y = y1 - 8 - 17 * len(week_lines)
                    for line in week_lines:
                        draw.text((x0 + padding, week_y), line, font=f_tiny, fill=(112, 120, 136))
                        week_y += 17

    if footer:
        draw.text((left_w, height - 28), footer, font=f_tiny, fill=(150, 156, 168))

    _watermark(draw, width, height)

    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    return buffer.getvalue()


# ===== 时间与周次 =====


def _parse_iso(text: str) -> date | None:
    matched = re.search(r'(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})', str(text))
    if not matched:
        return None
    try:
        return date(int(matched.group(1)), int(matched.group(2)), int(matched.group(3)))
    except ValueError:
        return None


def _week_of(term_start: str | None, day: date) -> int | None:
    start = _parse_iso(term_start) if term_start else None
    if not start:
        return None
    delta = (day - start).days
    if delta < 0:
        return 0
    return delta // 7 + 1


def _weekday_key(day: date) -> str:
    return WEEKDAY_KEYS[day.weekday()]


def _parse_weeks(text: str) -> list[int]:
    weeks: set[int] = set()
    for part in str(text).split(','):
        part = part.strip()
        if not part:
            continue
        if '-' in part:
            try:
                first, last = part.split('-', 1)
                weeks.update(range(int(first), int(last) + 1))
            except ValueError:
                continue
        else:
            try:
                weeks.add(int(part))
            except ValueError:
                continue
    return sorted(weeks)


def _period_block(period: str) -> int:
    matched = re.search(r'(\d+)', str(period))
    return _block_of(int(matched.group(1))) if matched else 0


# 每大节的上课时间（24 小时制）
PERIOD_START = {1: (8, 0), 3: (10, 0), 5: (14, 0), 7: (16, 0), 9: (19, 0)}
# 一大节的时长（两小节 + 课间）
PERIOD_MIN = 100


def _period_dt(day: date, period: str) -> datetime | None:
    base = _period_block(period)
    if base not in PERIOD_START:
        return None
    hour, minute = PERIOD_START[base]
    return datetime(day.year, day.month, day.day, hour, minute)


def _period_span(period: str) -> tuple[str, str] | None:
    """'3-4节' -> ('10:00', '11:40')。"""
    numbers = [int(item) for item in re.findall(r'\d+', str(period))]
    if not numbers:
        return None
    start = PERIOD_START.get(_block_of(numbers[0]))
    end = PERIOD_START.get(_block_of(numbers[-1]))
    if not start or not end:
        return None
    start_dt = datetime(2000, 1, 1, *start)
    end_dt = datetime(2000, 1, 1, *end) + timedelta(minutes=PERIOD_MIN)
    return f'{start_dt:%H:%M}', f'{end_dt:%H:%M}'


def _human_delta(target: datetime, now: datetime) -> str:
    minutes = int((target - now).total_seconds() // 60)
    if minutes < 60:
        return f'{minutes} 分钟后'
    hours, rest = divmod(minutes, 60)
    if hours < 24:
        return f'{hours} 小时 {rest} 分后' if rest else f'{hours} 小时后'
    days, hours = divmod(hours, 24)
    return f'{days} 天 {hours} 小时后' if hours else f'{days} 天后'


def next_lesson(data: dict[str, Any], after: datetime) -> tuple[datetime, dict[str, Any]] | None:
    """after 之后最近的一节课（最多往后找 21 天）。"""
    term = data.get('term_start') or DEFAULT_TERM_START
    for offset in range(0, 22):
        day = after.date() + timedelta(days=offset)
        week = _week_of(term, day)
        if not week or week < 1:
            continue
        key = _weekday_key(day)
        best: tuple[datetime, dict[str, Any]] | None = None
        for course in data.get('courses', []):
            if course.get('day') != key or week not in course.get('weeks', []):
                continue
            start = _period_dt(day, course.get('period', ''))
            if not start or start <= after:
                continue
            if best is None or start < best[0]:
                best = (start, course)
        if best:
            return best
    return None


# ===== 解析课表文件 =====


def _grid_from_xls(path: Path) -> list[list[str]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    rows: list[list[str]] = []
    for row_index in range(sheet.nrows):
        row: list[str] = []
        for col_index in range(sheet.ncols):
            value = sheet.cell(row_index, col_index).value
            if isinstance(value, float) and value == int(value):
                value = int(value)
            row.append('' if value is None else str(value))
        rows.append(row)
    return rows


def _grid_from_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True)
    worksheet = workbook.worksheets[0]
    return [['' if value is None else str(value) for value in row] for row in worksheet.iter_rows(values_only=True)]


def parse_schedule(path: Path) -> dict[str, Any]:
    """解析教务导出的课表文件，返回 {term, courses, updated}。"""
    grid = _grid_from_xlsx(path) if path.suffix.lower() in ('.xlsx', '.xlsm') else _grid_from_xls(path)

    header_text = '\n'.join(' '.join(cell for cell in row if cell) for row in grid[:4])
    term = ''
    matched = re.search(r'学年学期[：:]\s*([0-9\-]+)', header_text)
    if matched:
        term = matched.group(1)

    header_index = None
    day_columns: dict[int, str] = {}
    for index, row in enumerate(grid[:12]):
        hits = {
            col: weekday for col, cell in enumerate(row) for weekday in WEEKDAY_KEYS if weekday in str(cell)
        }
        if len(hits) >= 5:
            header_index = index
            day_columns = hits
            break
    if header_index is None:
        raise ValueError('没找到星期表头，可能不是教务导出的课表文件')

    courses: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    for row in grid[header_index + 1 :]:
        label = str(row[0]).strip() if row else ''
        period_guess = PERIOD_LABEL.get(label, '')
        for col, weekday in sorted(day_columns.items()):
            cell = str(row[col]) if col < len(row) else ''
            if not cell.strip():
                continue
            for matched in ENTRY_RE.finditer(cell):
                name = matched.group('name').strip()
                weeks = _parse_weeks(matched.group('weeks'))
                if not weeks:
                    continue
                period = matched.group('period').strip()
                location = (matched.group('loc') or '').strip()
                key = (name, weekday, tuple(weeks), period, location)
                if key in seen:
                    continue
                seen.add(key)
                courses.append(
                    {
                        'name': name,
                        'day': weekday,
                        'weeks': weeks,
                        'period': period,
                        'period_label': period_guess,
                        'location': location,
                    }
                )

    courses.sort(key=lambda item: (WEEKDAY_KEYS.index(item['day']), _period_block(item['period'])))
    return {'term': term, 'courses': courses, 'updated': int(time.time())}


# ===== 存取 =====


def _user_file(platform: str, user_id: str) -> Path:
    safe = re.sub(r'[^\w.-]', '_', f'{platform}_{user_id}')
    return DATA_DIR / f'{safe}.json'


def load_user(platform: str, user_id: str) -> dict[str, Any] | None:
    path = _user_file(platform, user_id)
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except Exception:
        return None


def save_user(platform: str, user_id: str, data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    _user_file(platform, user_id).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def iter_users():
    """遍历所有已导入课表的用户数据。"""
    if not DATA_DIR.exists():
        return
    for path in DATA_DIR.glob('*.json'):
        try:
            data = json.loads(path.read_text(encoding='utf-8'))
        except Exception:
            continue
        if data.get('courses'):
            yield path, data


# ===== 渲染入口 =====


def _week_dates(term_start: str | None, week: int) -> dict[str, str]:
    start = _parse_iso(term_start) if term_start else None
    if not start:
        return {}
    monday = start + timedelta(weeks=week - 1)
    return {WEEKDAY_KEYS[index]: f'{(monday + timedelta(days=index)):%m-%d}' for index in range(7)}


async def week_image(data: dict[str, Any], week: int) -> bytes:
    """整周课表图。"""
    courses = [course for course in data.get('courses', []) if week in course.get('weeks', [])]
    term = data.get('term') or '本学期'
    await ensure_font()
    return render_grid(
        courses,
        title=term,
        subtitle=f'第 {week} 周',
        show_weeks=False,
        footer='课表插件 · /周课表 · /今日课表 · /明日课表',
        day_dates=_week_dates(data.get('term_start'), week),
    )


async def day_image(data: dict[str, Any], target: date, label: str) -> bytes:
    """某一天的课表图（那一列高亮）。"""
    week = _week_of(data.get('term_start') or DEFAULT_TERM_START, target)
    day_key = _weekday_key(target)
    courses = [
        course
        for course in data.get('courses', [])
        if course.get('day') == day_key and (not week or week in course.get('weeks', []))
    ]
    subtitle = f'{label} {target:%m-%d} 周{WEEKDAYS[WEEKDAY_KEYS.index(day_key)][1]}'
    if week:
        subtitle += f' · 第 {week} 周'
    await ensure_font()
    return render_grid(
        courses,
        title=data.get('term') or '本学期',
        subtitle=subtitle,
        footer='课表插件 · /周课表 · /今日课表 · /明日课表',
        highlight=day_key,
        day_dates=_week_dates(data.get('term_start') or DEFAULT_TERM_START, week) if week else None,
    )


def next_lesson_line(data: dict[str, Any], now: datetime) -> str:
    found = next_lesson(data, now)
    if not found:
        return '往后 21 天都没课喵～'
    start, course = found
    span = _period_span(course.get('period', ''))
    clock = f'{span[0]}-{span[1]}' if span else f'{start:%H:%M}'
    detail = ' · '.join(part for part in (clock, course.get('name', ''), course.get('location', '')) if part)
    return f'下一节课：{start:%m-%d %H:%M}（{_human_delta(start, now)}）\n{detail}'


HELP_TEXT = (
    '课表用法：\n'
    '1. 把教务系统导出的「学生个人课表_*.xls / *.xlsx」发给我（群里发群文件、或私聊发文件都行）\n'
    '2. /课表            看本周课表（出图）\n'
    '3. /周课表 8        看第 8 周\n'
    '4. /今日课表        今天的课\n'
    '5. /明日课表        明天的课\n'
    '6. /下一节课        下一节是哪节课\n'
    '7. /上课提醒 开     课前 20 分钟提醒（默认关闭）\n'
    '8. /课表开学 2026-09-07  设置第 1 周的周一\n'
    '9. /课表清除        删掉已导入的课表\n'
    '\n'
    f'—— {WATERMARK}'
)

NO_DATA_TEXT = '还没导入课表喵～把教务导出的「学生个人课表_*.xls」发给我就行，发 /课表帮助 看用法。'

# ===== 指令 =====


def _platform_of(event: MessageEvent) -> str:
    return 'qq'


def _user_id_of(event: MessageEvent) -> str:
    return str(event.user_id)


def _plain(event: MessageEvent) -> str:
    return event.get_plaintext().strip()


schedule_cmd = on_command('课表', priority=0, block=True)
week_cmd = on_command('周课表', priority=0, block=True)
today_cmd = on_command('今日课表', priority=0, block=True)
tomorrow_cmd = on_command('明日课表', priority=0, block=True)
next_cmd = on_command('下一节课', priority=0, block=True)
help_cmd = on_command('课表帮助', priority=0, block=True)
term_cmd = on_command('课表开学', priority=0, block=True)
clear_cmd = on_command('课表清除', priority=0, block=True)
remind_cmd = on_command('上课提醒', priority=0, block=True)


def _week_from_args(text: str, data: dict[str, Any]) -> int | None:
    numbers = re.findall(r'\d+', text)
    if numbers:
        return int(numbers[0])
    return _week_of(data.get('term_start') or DEFAULT_TERM_START, datetime.now().date())


@schedule_cmd.handle()
async def handle_schedule(event: MessageEvent, args: Message = CommandArg()):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await schedule_cmd.finish(NO_DATA_TEXT)
    text = args.extract_plain_text().strip()
    if text in ('帮助', 'help', '?', '？'):
        await schedule_cmd.finish(HELP_TEXT)
    week = _week_from_args(text, data)
    if not week or week < 1:
        await schedule_cmd.finish('还没设置开学日期喵～用 /课表开学 2026-09-07 设一下。')
    await schedule_cmd.finish(MessageSegment.image(await week_image(data, week)))


@week_cmd.handle()
async def handle_week(event: MessageEvent, args: Message = CommandArg()):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await week_cmd.finish(NO_DATA_TEXT)
    week = _week_from_args(args.extract_plain_text().strip(), data)
    if not week or week < 1:
        await week_cmd.finish('还没设置开学日期喵～用 /课表开学 2026-09-07 设一下。')
    await week_cmd.finish(MessageSegment.image(await week_image(data, week)))


@today_cmd.handle()
async def handle_today(event: MessageEvent):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await today_cmd.finish(NO_DATA_TEXT)
    today = datetime.now().date()
    await today_cmd.finish(MessageSegment.image(await day_image(data, today, '今天')))


@tomorrow_cmd.handle()
async def handle_tomorrow(event: MessageEvent):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await tomorrow_cmd.finish(NO_DATA_TEXT)
    tomorrow = datetime.now().date() + timedelta(days=1)
    await tomorrow_cmd.finish(MessageSegment.image(await day_image(data, tomorrow, '明天')))


@next_cmd.handle()
async def handle_next(event: MessageEvent):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await next_cmd.finish(NO_DATA_TEXT)
    await next_cmd.finish(next_lesson_line(data, datetime.now()))


@help_cmd.handle()
async def handle_help():
    await help_cmd.finish(HELP_TEXT)


@term_cmd.handle()
async def handle_term(event: MessageEvent, args: Message = CommandArg()):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await term_cmd.finish(NO_DATA_TEXT)
    target = _parse_iso(args.extract_plain_text())
    if not target:
        await term_cmd.finish('用法：/课表开学 2026-09-07（第 1 周的周一）')
    data['term_start'] = f'{target:%Y-%m-%d}'
    save_user(platform, user_id, data)
    await term_cmd.finish(f'好，第 1 周从 {target:%Y-%m-%d}（周{WEEKDAYS[target.weekday()][1]}）开始算。')


@clear_cmd.handle()
async def handle_clear(event: MessageEvent):
    platform, user_id = _platform_of(event), _user_id_of(event)
    path = _user_file(platform, user_id)
    if not path.exists():
        await clear_cmd.finish(NO_DATA_TEXT)
    path.unlink(missing_ok=True)
    _sent.clear()
    await clear_cmd.finish('课表已删除，重新发一份文件就能再导入。')


@remind_cmd.handle()
async def handle_remind(event: MessageEvent, args: Message = CommandArg()):
    platform, user_id = _platform_of(event), _user_id_of(event)
    data = load_user(platform, user_id)
    if not data:
        await remind_cmd.finish(NO_DATA_TEXT)
    text = args.extract_plain_text().strip()
    if text in ('开', '开启', 'on', '打开'):
        data['reminder'] = True
        save_user(platform, user_id, data)
        detail = '群里会 @ 你' if event.message_type == 'group' else '发私聊给你'
        await remind_cmd.finish(f'上课提醒已开启：每节课前 {REMIND_MIN} 分钟{detail}。')
        return
    if text in ('关', '关闭', 'off', '关掉'):
        data['reminder'] = False
        save_user(platform, user_id, data)
        await remind_cmd.finish('上课提醒已关闭。')
        return
    await remind_cmd.finish(
        f'上课提醒当前是「{"开" if data.get("reminder") else "关"}」的。\n'
        '用 /上课提醒 开  或  /上课提醒 关  切换。'
    )


# ===== 收文件 =====


async def _download(url: str) -> bytes | None:
    async with AsyncClient(timeout=30, follow_redirects=True) as client:
        response = await client.get(url)
        if response.status_code != 200:
            return None
        content = response.content
    if not content or len(content) > FILE_MAX_BYTES:
        return None
    return content


async def _store_schedule(event: MessageEvent, bot: Bot, filename: str, content: bytes) -> str:
    """解析上传的课表并保存，返回给用户的回执文本。"""
    platform, user_id = _platform_of(event), _user_id_of(event)
    suffix = Path(filename or 'schedule.xls').suffix.lower()
    if suffix not in ('.xls', '.xlsx', '.xlsm'):
        return '只认 .xls / .xlsx 的课表文件喵～'

    tmp = DATA_DIR / f'_upload_{int(time.time() * 1000)}{suffix}'
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        tmp.write_bytes(content)
        parsed = parse_schedule(tmp)
    except Exception as error:
        return f'课表解析失败：{str(error)[:60]}'
    finally:
        tmp.unlink(missing_ok=True)

    if not parsed['courses']:
        return '没解析出课程，确认一下是教务导出的「学生个人课表」文件喵～'

    data = {
        'user_id': user_id,
        'platform': platform,
        'message_type': event.message_type,
        'group_id': str(event.group_id) if getattr(event, 'group_id', None) else '',
        'term': parsed['term'],
        'term_start': DEFAULT_TERM_START,
        'courses': parsed['courses'],
        'reminder': False,
        'updated': int(time.time()),
    }
    save_user(platform, user_id, data)
    _sent.clear()

    days = len({course['day'] for course in data['courses']})
    return (
        f'课表已导入：{parsed["term"] or "本学期"} · {len(data["courses"])} 节课 · 分布在 {days} 天\n'
        f'用 /课表 看本周，/今日课表 看今天，/上课提醒 开 打开课前 {REMIND_MIN} 分钟提醒。'
    )


group_file_notice = on_notice(priority=0, block=True, rule=lambda event: isinstance(event, GroupUploadNoticeEvent))
file_message = on_message(priority=0, block=True, rule=lambda event: any(segment.type == 'file' for segment in event.get_message()))


@group_file_notice.handle()
async def handle_group_file(event: GroupUploadNoticeEvent, bot: Bot):
    """群里上传群文件时触发。"""
    file_info = event.file
    if file_info.size and int(file_info.size) > FILE_MAX_BYTES:
        await group_file_notice.finish('文件太大了喵～课表一般在 100KB 以内。')
    url = str(getattr(file_info, 'url', '') or '').strip()
    if not url:
        try:
            result = await bot.call_api('get_group_file_url', group_id=event.group_id, file_id=file_info.id)
            url = str((result or {}).get('url') or '').strip()
        except Exception:
            url = ''
    if not url:
        await group_file_notice.finish('拿不到这个文件的下载地址，试试私聊把课表文件发给我。')
    content = await _download(url)
    if content is None:
        await group_file_notice.finish('文件下载失败或超过 5MB，再传一次试试喵～')
    await group_file_notice.finish(await _store_schedule(event, bot, file_info.name, content))


@file_message.handle()
async def handle_file_message(event: MessageEvent, bot: Bot):
    """私聊 / 群里直接发文件时触发。"""
    segment = next((item for item in event.get_message() if item.type == 'file'), None)
    if segment is None:
        return
    url = str(segment.data.get('url') or '').strip()
    name = str(segment.data.get('file') or segment.data.get('name') or 'schedule.xls')
    if not url:
        file_id = str(segment.data.get('file_id') or '').strip()
        if file_id:
            try:
                result = await bot.call_api('get_file', file_id=file_id)
                url = str((result or {}).get('url') or '').strip()
            except Exception:
                url = ''
    if not url:
        await file_message.finish('拿不到这个文件的下载地址，试试把课表当「群文件」上传喵～')
    content = await _download(url)
    if content is None:
        await file_message.finish('文件下载失败或超过 5MB，再传一次试试喵～')
    await file_message.finish(await _store_schedule(event, bot, name, content))


# ===== 上课提醒 =====

# 已经提醒过的课：{(用户键, 日期, 节次, 课名)}
_sent: dict[str, set[str]] = {}
_reminder_task: asyncio.Task | None = None


def _target_of(data: dict[str, Any]) -> tuple[str, int] | None:
    """把存下来的目标还原成 (group/private, id)。"""
    if data.get('message_type') == 'group' and data.get('group_id'):
        return 'group', int(data['group_id'])
    user_id = str(data.get('user_id') or '').strip()
    if user_id.isdigit():
        return 'private', int(user_id)
    return None


async def _tick_reminders() -> None:
    now = datetime.now()
    for path, data in iter_users():
        if not data.get('reminder'):
            continue
        target = _target_of(data)
        if target is None:
            continue
        for course in data.get('courses', []):
            if course.get('day') != _weekday_key(now.date()):
                continue
            week = _week_of(data.get('term_start') or DEFAULT_TERM_START, now.date())
            if not week or week not in course.get('weeks', []):
                continue
            start = _period_dt(now.date(), course.get('period', ''))
            if not start:
                continue
            lead = (start - now).total_seconds() / 60
            if not (REMIND_MIN - 0.5 <= lead <= REMIND_MIN + 0.5):
                continue
            tag = f'{path.name}|{now:%Y-%m-%d}|{course.get("period")}|{course.get("name")}'
            if tag in _sent.get(path.name, set()):
                continue
            span = _period_span(course.get('period', ''))
            clock = f'{span[0]}-{span[1]}' if span else f'{start:%H:%M}'
            detail = ' · '.join(part for part in (clock, course.get('name', ''), course.get('location', '')) if part)
            message = Message(f'⏰ 还有 {REMIND_MIN} 分钟上课\n{detail}')
            kind, target_id = target
            try:
                if kind == 'group':
                    await _send_group(target_id, data.get('user_id'), message)
                else:
                    await _send_private(target_id, message)
            except Exception:
                continue
            _sent.setdefault(path.name, set()).add(tag)


async def _send_group(group_id: int, user_id: str | None, message: Message) -> None:
    from nonebot import get_bots

    for bot in get_bots().values():
        payload = message
        if user_id and str(user_id).isdigit():
            payload = MessageSegment.at(int(user_id)) + Message('\n') + message
        await bot.send_group_msg(group_id=group_id, message=payload)
        return


async def _send_private(user_id: int, message: Message) -> None:
    from nonebot import get_bots

    for bot in get_bots().values():
        await bot.send_private_msg(user_id=user_id, message=message)
        return


async def _reminder_loop() -> None:
    while True:
        try:
            await _tick_reminders()
        except Exception:
            pass
        await asyncio.sleep(REMIND_TICK)


driver = get_driver()


@driver.on_startup
async def _start_reminder_loop() -> None:
    global _reminder_task
    if _reminder_task is None or _reminder_task.done():
        _reminder_task = asyncio.create_task(_reminder_loop())


@driver.on_shutdown
async def _stop_reminder_loop() -> None:
    global _reminder_task
    if _reminder_task is not None:
        _reminder_task.cancel()
        _reminder_task = None


__plugin_meta__ = PluginMetadata(
    name='课表查询',
    description='导入教务课表、按周出图，课前提醒（提醒默认关闭）',
    usage='/课表 与 /课表帮助，例如：/课表、/周课表 8、/今日课表、/明日课表、/上课提醒 开',
    type='application',
    extra={'author': WATERMARK, 'version': '1.0'},
)

register_help_commands(
    [
        {'name': '课表', 'description': '导入并查看课表（出图）', 'usage': '/课表 [周次]'},
        {'name': '周课表', 'description': '查看指定周次的课表', 'usage': '/周课表 <周次>'},
        {'name': '今日课表', 'description': '今天的课表', 'usage': '/今日课表'},
        {'name': '明日课表', 'description': '明天的课表', 'usage': '/明日课表'},
        {'name': '下一节课', 'description': '下一节是什么课', 'usage': '/下一节课'},
        {'name': '上课提醒', 'description': '上课前 20 分钟提醒（默认关闭）', 'usage': '/上课提醒 开|关'},
        {'name': '课表开学', 'description': '设置第 1 周的周一', 'usage': '/课表开学 2026-09-07'},
        {'name': '课表清除', 'description': '删除已导入的课表', 'usage': '/课表清除'},
        {'name': '课表帮助', 'description': '课表插件用法', 'usage': '/课表帮助'},
    ]
)
